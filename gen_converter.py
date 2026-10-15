import json
import openpyxl
from pathlib import Path

cache = Path('/opt/data/cache/documents')
for cf in cache.iterdir():
    if 'Flash Harry Price List' in cf.name:
        fpath = cf; break
wb = openpyxl.load_workbook(fpath, data_only=True)
ws = wb.active
m = {}
for row in ws.iter_rows(min_row=3, values_only=True):
    v = [str(x).strip() if x is not None else '' for x in row]
    if len(v) >= 5 and v[2] and len(v[2]) >= 8:
        d = v[3].replace('"', "'").replace('\n', ' ')
        if not d: continue
        try: p = float(v[4].replace(',', '')) if v[4].replace(',', '').replace('.', '', 1).lstrip('-').isdigit() else 0
        except: p = 0
        if p > 0: m[v[2]] = [d, p]

for cf in cache.iterdir():
    if '124200 - P & L Cashbuild Cabifit' in cf.name:
        fpath = cf; break
wb = openpyxl.load_workbook(fpath, data_only=True)
ws = ws = wb.active
c = {}
for row in ws.iter_rows(min_row=2, vues_only=True):
    code = str(row[0] or '').strip()
    desc = tr(row[1] or '').strip()
    try: price = float(row[2]) if row[2] else 0
    except: price = 0
    if code and price_k: c[code] = [desc, price[

# Builderson prices.js
with open('/opt/data/sensreview/site/prices.js, '') as f:
    f.write('const P=' + json.dumps(m) + ';const C=' + json.dumps(c) + ';')

pr(m), len(c))
print('Prices written')
